from __future__ import annotations

import importlib.metadata
import io
from typing import Any

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from codegate_api.documents.models import (
    CreationPayload,
    DocumentFormat,
    DocumentOperation,
    MarkdownCreationPayload,
    NativeLocator,
    StructuralDiff,
    StructureItem,
    TableCellSetOperation,
    TextReplaceOperation,
)
from codegate_api.documents.writers.base import ProposedDocument, WriterError
from codegate_api.documents.writers.text import markdown_blocks, replace_across_runs


class PptxWriter:
    format = DocumentFormat.PPTX

    def __init__(self) -> None:
        self.fingerprint = "python-pptx/" + importlib.metadata.version("python-pptx")

    def read_structure(self, source: bytes) -> list[StructureItem]:
        presentation = _load(source)
        items: list[StructureItem] = []
        for slide_index, slide in enumerate(presentation.slides):
            for shape_index, shape in enumerate(slide.shapes):
                if getattr(shape, "has_text_frame", False):
                    for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
                        items.append(
                            StructureItem(
                                locator=NativeLocator(
                                    kind="slide_paragraph",
                                    slide_index=slide_index,
                                    shape_index=shape_index,
                                    paragraph_index=paragraph_index,
                                ),
                                value=paragraph.text,
                                value_type="string",
                            )
                        )
                if getattr(shape, "has_table", False):
                    for row_index, row in enumerate(shape.table.rows):
                        for column_index, cell in enumerate(row.cells):
                            items.append(
                                StructureItem(
                                    locator=NativeLocator(
                                        kind="slide_table_cell",
                                        slide_index=slide_index,
                                        shape_index=shape_index,
                                        row_index=row_index,
                                        column_index=column_index,
                                    ),
                                    value=cell.text,
                                    value_type="string",
                                )
                            )
        return items

    def mutate(
        self,
        source: bytes,
        operations: list[DocumentOperation],
    ) -> ProposedDocument:
        presentation = _load(source)
        diffs: list[StructuralDiff] = []
        for index, operation in enumerate(operations):
            if isinstance(operation, TextReplaceOperation):
                paragraph = self._paragraph(presentation, operation.locator)
                before = paragraph.text
                if not paragraph.runs:
                    paragraph.add_run().text = before
                replace_across_runs(paragraph.runs, operation.expected, operation.replacement)
                after = paragraph.text
            elif isinstance(operation, TableCellSetOperation):
                cell = self._cell(presentation, operation.locator)
                before = cell.text
                paragraph = cell.text_frame.paragraphs[0]
                if not paragraph.runs:
                    paragraph.add_run().text = before
                replace_across_runs(paragraph.runs, operation.expected, operation.replacement)
                after = cell.text
            else:
                raise WriterError("unsupported_operation", "PPTX operation is unsupported")
            diffs.append(
                StructuralDiff(
                    operation_index=index,
                    operation_type=operation.type,
                    locator=operation.locator,
                    before=before,
                    after=after,
                )
            )
        result = io.BytesIO()
        presentation.save(result)
        proposed = result.getvalue()
        _load(proposed)
        return ProposedDocument(proposed, diffs, [], self.fingerprint)

    def create(
        self,
        payload: CreationPayload,
        *,
        template: bytes | None = None,
    ) -> ProposedDocument:
        if not isinstance(payload, MarkdownCreationPayload):
            raise WriterError("unsupported_operation", "PPTX creation requires Markdown")
        presentation = (
            Presentation(io.BytesIO(template)) if template is not None else Presentation()
        )
        blocks = markdown_blocks(payload.markdown)
        slides: list[tuple[str, list[tuple[str, str]]]] = []
        title = "Document"
        body: list[tuple[str, str]] = []
        for kind, value in blocks:
            assert isinstance(value, str)
            if kind in {"heading1", "heading2"}:
                if body or title != "Document":
                    slides.append((title, body))
                title, body = value, []
            else:
                body.append((kind, value))
        slides.append((title, body))
        for title, body in slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title
            placeholder = slide.placeholders[1]
            text_frame = placeholder.text_frame
            text_frame.clear()
            for item_index, (kind, value) in enumerate(body):
                paragraph = (
                    text_frame.paragraphs[0] if item_index == 0 else text_frame.add_paragraph()
                )
                paragraph.text = value
                paragraph.level = 1 if kind == "bullet" else 0
                paragraph.alignment = PP_ALIGN.LEFT
            if not body:
                placeholder.left = Inches(1)
        result = io.BytesIO()
        presentation.save(result)
        proposed = result.getvalue()
        _load(proposed)
        fingerprint = self.fingerprint + ("+approved-template" if template is not None else "")
        return ProposedDocument(proposed, [], [], fingerprint)

    @staticmethod
    def _shape(presentation: Any, locator: NativeLocator) -> Any:
        if locator.slide_index is None or locator.shape_index is None:
            raise WriterError("invalid_locator", "PPTX locator requires slide and shape")
        try:
            return presentation.slides[locator.slide_index].shapes[locator.shape_index]
        except IndexError as error:
            raise WriterError("invalid_locator", "PPTX shape does not exist") from error

    def _paragraph(self, presentation: Any, locator: NativeLocator) -> Any:
        if locator.kind != "slide_paragraph" or locator.paragraph_index is None:
            raise WriterError("invalid_locator", "PPTX text requires a paragraph locator")
        shape = self._shape(presentation, locator)
        if not getattr(shape, "has_text_frame", False):
            raise WriterError("invalid_locator", "PPTX shape has no text frame")
        try:
            return shape.text_frame.paragraphs[locator.paragraph_index]
        except IndexError as error:
            raise WriterError("invalid_locator", "PPTX paragraph does not exist") from error

    def _cell(self, presentation: Any, locator: NativeLocator) -> Any:
        if (
            locator.kind != "slide_table_cell"
            or locator.row_index is None
            or locator.column_index is None
        ):
            raise WriterError("invalid_locator", "PPTX table operation requires a cell locator")
        shape = self._shape(presentation, locator)
        if not getattr(shape, "has_table", False):
            raise WriterError("invalid_locator", "PPTX shape has no table")
        try:
            return shape.table.cell(locator.row_index, locator.column_index)
        except IndexError as error:
            raise WriterError("invalid_locator", "PPTX table cell does not exist") from error


def _load(source: bytes) -> Any:
    try:
        return Presentation(io.BytesIO(source))
    except Exception as error:
        raise WriterError("malformed_document", "PPTX could not be parsed") from error
