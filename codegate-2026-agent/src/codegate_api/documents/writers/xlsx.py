from __future__ import annotations

import importlib.metadata
import io
from typing import Any

from openpyxl import Workbook, load_workbook

from codegate_api.documents.models import (
    BooleanValue,
    CellValue,
    CreationPayload,
    DocumentFormat,
    DocumentOperation,
    FormulaValue,
    NativeLocator,
    NullValue,
    NumberValue,
    SpreadsheetCellsSetOperation,
    StringValue,
    StructuralDiff,
    StructureItem,
    WorkbookCreationPayload,
)
from codegate_api.documents.writers.base import ProposedDocument, WriterError


class XlsxWriter:
    format = DocumentFormat.XLSX

    def __init__(self) -> None:
        self.fingerprint = "openpyxl/" + importlib.metadata.version("openpyxl")

    def read_structure(self, source: bytes) -> list[StructureItem]:
        workbook = _load(source, data_only=False)
        items: list[StructureItem] = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    items.append(
                        StructureItem(
                            locator=NativeLocator(
                                kind="spreadsheet_cell",
                                sheet_name=worksheet.title,
                                address=cell.coordinate,
                            ),
                            value=_value_model(cell.value, cell.data_type).model_dump(mode="json"),
                            value_type=cell.data_type,
                        )
                    )
        return items

    def mutate(
        self,
        source: bytes,
        operations: list[DocumentOperation],
    ) -> ProposedDocument:
        workbook = _load(source, data_only=False)
        diffs: list[StructuralDiff] = []
        operation_index = 0
        for operation in operations:
            if not isinstance(operation, SpreadsheetCellsSetOperation):
                raise WriterError("unsupported_operation", "XLSX operation is unsupported")
            for change in operation.cells:
                if change.sheet_name not in workbook.sheetnames:
                    raise WriterError("invalid_locator", "XLSX sheet does not exist")
                cell = workbook[change.sheet_name][change.address]
                actual = _value_model(cell.value, cell.data_type)
                if actual != change.expected:
                    raise WriterError(
                        "expected_value_mismatch",
                        f"XLSX expected value mismatch at {change.sheet_name}!{change.address}",
                    )
                _set_cell(cell, change.replacement)
                diffs.append(
                    StructuralDiff(
                        operation_index=operation_index,
                        operation_type=operation.type,
                        locator=NativeLocator(
                            kind="spreadsheet_cell",
                            sheet_name=change.sheet_name,
                            address=change.address,
                        ),
                        before=change.expected.model_dump(mode="json"),
                        after=change.replacement.model_dump(mode="json"),
                    )
                )
                operation_index += 1
        result = io.BytesIO()
        workbook.save(result)
        proposed = result.getvalue()
        _load(proposed, data_only=False)
        return ProposedDocument(proposed, diffs, [], self.fingerprint)

    def create(
        self,
        payload: CreationPayload,
        *,
        template: bytes | None = None,
    ) -> ProposedDocument:
        if template is not None:
            raise WriterError("unsupported_template", "XLSX templates are not supported")
        if not isinstance(payload, WorkbookCreationPayload):
            raise WriterError("unsupported_operation", "XLSX creation requires typed sheets")
        workbook = Workbook()
        if workbook.active is not None:
            workbook.remove(workbook.active)
        for sheet in payload.sheets:
            if sheet.name in workbook.sheetnames:
                raise WriterError("duplicate_sheet", f"duplicate sheet: {sheet.name}")
            worksheet = workbook.create_sheet(sheet.name)
            seen: set[str] = set()
            for item in sheet.cells:
                if item.address in seen:
                    raise WriterError(
                        "duplicate_cell", f"duplicate cell: {sheet.name}!{item.address}"
                    )
                seen.add(item.address)
                _set_cell(worksheet[item.address], item.value)
        result = io.BytesIO()
        workbook.save(result)
        proposed = result.getvalue()
        _load(proposed, data_only=False)
        return ProposedDocument(proposed, [], [], self.fingerprint)


def _load(source: bytes, *, data_only: bool) -> Any:
    try:
        return load_workbook(io.BytesIO(source), data_only=data_only, keep_links=False)
    except Exception as error:
        raise WriterError("malformed_document", "XLSX could not be parsed") from error


def _value_model(value: Any, data_type: str) -> CellValue:
    if value is None:
        return NullValue()
    if data_type == "f":
        formula = str(value)
        return FormulaValue(value=formula if formula.startswith("=") else "=" + formula)
    if isinstance(value, bool):
        return BooleanValue(value=value)
    if isinstance(value, int | float):
        return NumberValue(value=float(value))
    return StringValue(value=str(value))


def _set_cell(cell: Any, value: CellValue) -> None:
    if isinstance(value, NullValue):
        cell.value = None
    elif isinstance(value, FormulaValue):
        cell.value = value.value
        cell.data_type = "f"
    elif isinstance(value, StringValue):
        cell.value = value.value
        cell.data_type = "s"
    else:
        cell.value = value.value
