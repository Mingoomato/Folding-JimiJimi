from __future__ import annotations

import importlib.metadata
import io
from typing import Any

from docx import Document

from codegate_api.documents.models import (
    CreationPayload,
    DocumentFormat,
    DocumentOperation,
    MarkdownCreationPayload,
    NativeLocator,
    StructuralDiff,
    StructureItem,
    TableCellSetOperation,
    TextReplaceOperation,
)
from codegate_api.documents.writers.base import ProposedDocument, WriterError
from codegate_api.documents.writers.text import markdown_blocks, replace_across_runs


class DocxWriter:
    format = DocumentFormat.DOCX

    def __init__(self) -> None:
        self.fingerprint = "python-docx/" + importlib.metadata.version("python-docx")

    def read_structure(self, source: bytes) -> list[StructureItem]:
        document = _load(source)
        items = [
            StructureItem(
                locator=NativeLocator(kind="paragraph", block_index=index),
                value=paragraph.text,
                value_type="string",
            )
            for index, paragraph in enumerate(document.paragraphs)
        ]
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    items.append(
                        StructureItem(
                            locator=NativeLocator(
                                kind="table_cell",
                                table_index=table_index,
                                row_index=row_index,
                                column_index=column_index,
                            ),
                            value=cell.text,
                            value_type="string",
                        )
                    )
        return items

    def mutate(
        self,
        source: bytes,
        operations: list[DocumentOperation],
    ) -> ProposedDocument:
        document = _load(source)
        diffs: list[StructuralDiff] = []
        for index, operation in enumerate(operations):
            if isinstance(operation, TextReplaceOperation):
                paragraph = self._paragraph(document, operation.locator)
                before = paragraph.text
                replace_across_runs(paragraph.runs, operation.expected, operation.replacement)
                after = paragraph.text
            elif isinstance(operation, TableCellSetOperation):
                cell = self._cell(document, operation.locator)
                before = cell.text
                paragraph = cell.add_paragraph() if not cell.paragraphs else cell.paragraphs[0]
                if not paragraph.runs:
                    paragraph.add_run(before)
                replace_across_runs(paragraph.runs, operation.expected, operation.replacement)
                after = cell.text
            else:
                raise WriterError("unsupported_operation", "DOCX operation is unsupported")
            diffs.append(
                StructuralDiff(
                    operation_index=index,
                    operation_type=operation.type,
                    locator=operation.locator,
                    before=before,
                    after=after,
                )
            )
        result = io.BytesIO()
        document.save(result)
        proposed = result.getvalue()
        _load(proposed)
        return ProposedDocument(proposed, diffs, [], self.fingerprint)

    def create(
        self,
        payload: CreationPayload,
        *,
        template: bytes | None = None,
    ) -> ProposedDocument:
        if not isinstance(payload, MarkdownCreationPayload):
            raise WriterError("unsupported_operation", "DOCX creation requires Markdown")
        document = Document(io.BytesIO(template)) if template is not None else Document()
        for kind, value in markdown_blocks(payload.markdown):
            assert isinstance(value, str)
            if kind.startswith("heading"):
                document.add_heading(value, level=min(int(kind[-1]), 9))
            elif kind == "bullet":
                document.add_paragraph(value, style="List Bullet")
            else:
                document.add_paragraph(value)
        result = io.BytesIO()
        document.save(result)
        proposed = result.getvalue()
        _load(proposed)
        fingerprint = self.fingerprint + ("+approved-template" if template is not None else "")
        return ProposedDocument(proposed, [], [], fingerprint)

    @staticmethod
    def _paragraph(document: Any, locator: NativeLocator) -> Any:
        if locator.kind != "paragraph" or locator.block_index is None:
            raise WriterError("invalid_locator", "DOCX text requires a paragraph locator")
        try:
            paragraph = document.paragraphs[locator.block_index]
        except IndexError as error:
            raise WriterError("invalid_locator", "DOCX paragraph does not exist") from error
        if not paragraph.runs:
            paragraph.add_run(paragraph.text)
        return paragraph

    @staticmethod
    def _cell(document: Any, locator: NativeLocator) -> Any:
        if (
            locator.kind != "table_cell"
            or locator.table_index is None
            or locator.row_index is None
            or locator.column_index is None
        ):
            raise WriterError("invalid_locator", "DOCX table operation requires a cell locator")
        try:
            return document.tables[locator.table_index].cell(
                locator.row_index,
                locator.column_index,
            )
        except IndexError as error:
            raise WriterError("invalid_locator", "DOCX table cell does not exist") from error


def _load(source: bytes) -> Any:
    try:
        return Document(io.BytesIO(source))
    except Exception as error:
        raise WriterError("malformed_document", "DOCX could not be parsed") from error
